from __future__ import annotations

import csv
import hashlib
import json
import re
import shutil
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Iterable
from uuid import uuid4

from openpyxl import load_workbook

from database import db
from vendor_models import VendorKnowledgeItem, VendorKnowledgeSource


HEADER_SCAN_ROWS = 40
SUPPORTED_EXTENSIONS = {".xlsx", ".xlsm", ".csv"}

_HEADER_ALIASES = {
    "description": {"item description", "description", "product description", "product", "item"},
    "part_number": {"part no", "part no.", "part number", "part #", "sku", "mpn", "manufacturer part number"},
    "quantity": {"qty", "quantity", "stock", "stock qty", "available", "availability"},
    "price_small": {"price qty 1- 29", "price qty 1-29", "price qty 1 - 29", "price 1-29", "price 1 - 29"},
    "price_large": {"price qty 30+", "price qty 30 +", "price 30+", "price 30 +"},
    "price": {"price", "unit price", "sell price", "dealer price"},
    "condition": {"condition", "grade", "item condition"},
    "vendor_item_id": {"id#", "id #", "item id", "stock id", "id"},
}

_STOPWORDS = {
    "and", "the", "for", "with", "from", "this", "that", "unit", "units", "item", "items",
    "required", "requirement", "requirements", "tender", "please", "find", "show", "model", "models",
    "business", "computer", "laptop", "desktop", "new", "stock", "price", "qty", "quantity",
}


@dataclass(frozen=True)
class ParsedVendorItem:
    source_sheet: str
    source_row: int
    source_category: str | None
    stock_type: str
    item_marker: str | None
    vendor_item_id: str | None
    description: str
    part_number: str | None
    quantity_text: str | None
    quantity_minimum: int | None
    quantity_exact: bool
    price_qty_1_29: Decimal | None
    price_qty_30_plus: Decimal | None
    unit_price: Decimal | None
    condition: str | None
    availability_status: str
    is_available: bool
    raw: dict


@dataclass(frozen=True)
class ParsedVendorDocument:
    detected_vendor: str | None
    source_date: date | None
    row_count: int
    items: list[ParsedVendorItem]


def normalize_vendor_name(value: str | None) -> str:
    cleaned = re.sub(r"[^a-z0-9]+", " ", (value or "").lower())
    return " ".join(cleaned.split())


def safe_slug(value: str) -> str:
    slug = re.sub(r"[^A-Za-z0-9._-]+", "-", value.strip()).strip("-._")
    return slug[:80] or "vendor"


def sha256_file(path: str | Path) -> str:
    digest = hashlib.sha256()
    with Path(path).open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _clean_text(value) -> str:
    if value is None:
        return ""
    return " ".join(str(value).replace("\xa0", " ").split()).strip()


def _header_text(value) -> str:
    text = _clean_text(value).lower()
    text = re.sub(r"\s+", " ", text)
    return text.strip(" :")


def _find_header_map(rows: list[list], max_rows: int = HEADER_SCAN_ROWS) -> tuple[int, dict[str, int]] | None:
    for row_index, row in enumerate(rows[:max_rows], start=1):
        normalized = [_header_text(value) for value in row]
        mapping: dict[str, int] = {}
        for field, aliases in _HEADER_ALIASES.items():
            for col_index, value in enumerate(normalized):
                if value in aliases:
                    mapping[field] = col_index
                    break
        if "description" in mapping and "part_number" in mapping:
            return row_index, mapping
    return None


def _parse_decimal(value) -> Decimal | None:
    if value in (None, ""):
        return None
    if isinstance(value, Decimal):
        return value
    if isinstance(value, (int, float)):
        return Decimal(str(value))
    text = _clean_text(value)
    if not text:
        return None
    text = re.sub(r"[^0-9.\-]", "", text)
    if not text or text in {"-", ".", "-."}:
        return None
    try:
        return Decimal(text)
    except InvalidOperation:
        return None


def _parse_quantity(value) -> tuple[str | None, int | None, bool, str, bool]:
    text = _clean_text(value)
    if not text:
        return None, None, True, "unknown", False
    lowered = text.lower()
    if lowered in {"n/a", "na", "none", "out", "out of stock", "sold", "sold out"}:
        return text, 0, True, "unavailable", False
    if lowered in {"in stock", "available", "yes", "y"}:
        return text, None, False, "available", True
    match = re.search(r"-?\d+", text.replace(",", ""))
    if not match:
        return text, None, False, "unknown", False
    quantity = max(0, int(match.group(0)))
    exact = "+" not in text and "more" not in lowered and ">" not in text
    available = quantity > 0
    return text, quantity, exact, "available" if available else "unavailable", available


def _stock_type_from_sheet(sheet_name: str) -> str:
    lowered = sheet_name.lower()
    if "open box" in lowered:
        return "open_box"
    if "refurb" in lowered:
        return "refurbished"
    return "new"


def _parse_date_from_text(text: str) -> date | None:
    match = re.search(r"\b(\d{1,2})(?:st|nd|rd|th)?\s+([A-Za-z]+)\s+(20\d{2})\b", text, flags=re.I)
    if not match:
        return None
    clean = f"{match.group(1)} {match.group(2)} {match.group(3)}"
    try:
        return datetime.strptime(clean, "%d %B %Y").date()
    except ValueError:
        try:
            return datetime.strptime(clean, "%d %b %Y").date()
        except ValueError:
            return None


def _detect_vendor_from_title(text: str) -> str | None:
    text = _clean_text(text)
    if not text:
        return None
    match = re.match(r"(.+?)\s+(?:refurbished\s+)?offers\b", text, flags=re.I)
    if match:
        return _clean_text(match.group(1))[:255]
    return None


def _category_candidate(row: list, mapping: dict[str, int]) -> str | None:
    desc_index = mapping["description"]
    description = _clean_text(row[desc_index] if desc_index < len(row) else None)
    if not description:
        return None
    part_index = mapping.get("part_number")
    part_number = _clean_text(row[part_index] if part_index is not None and part_index < len(row) else None)
    if part_number:
        return None
    populated_commercial = False
    for field in ("quantity", "price_small", "price_large", "price", "condition", "vendor_item_id"):
        idx = mapping.get(field)
        if idx is not None and idx < len(row) and _clean_text(row[idx]):
            populated_commercial = True
            break
    if populated_commercial:
        return None
    if len(description) > 160:
        return None
    return description.strip(" :")


def _build_item(sheet_name: str, row_number: int, row: list, mapping: dict[str, int], category: str | None) -> ParsedVendorItem | None:
    def value(field: str):
        idx = mapping.get(field)
        return row[idx] if idx is not None and idx < len(row) else None

    description = _clean_text(value("description"))
    part_number = _clean_text(value("part_number"))
    if not description or not part_number:
        return None
    if _header_text(description) in _HEADER_ALIASES["description"] and _header_text(part_number) in _HEADER_ALIASES["part_number"]:
        return None

    quantity_text, quantity_minimum, quantity_exact, availability_status, is_available = _parse_quantity(value("quantity"))
    price_small = _parse_decimal(value("price_small"))
    price_large = _parse_decimal(value("price_large"))
    unit_price = _parse_decimal(value("price"))
    if unit_price is None:
        unit_price = price_small
    condition = _clean_text(value("condition")) or None
    vendor_item_id = _clean_text(value("vendor_item_id")) or None

    item_marker = None
    first_non_empty = next((_clean_text(cell) for cell in row if _clean_text(cell)), "")
    if first_non_empty and first_non_empty not in {description, part_number, vendor_item_id} and len(first_non_empty) <= 100:
        item_marker = first_non_empty

    raw = {f"column_{index + 1}": cell for index, cell in enumerate(row) if cell not in (None, "")}
    return ParsedVendorItem(
        source_sheet=sheet_name,
        source_row=row_number,
        source_category=category,
        stock_type=_stock_type_from_sheet(sheet_name),
        item_marker=item_marker,
        vendor_item_id=vendor_item_id,
        description=description,
        part_number=part_number,
        quantity_text=quantity_text,
        quantity_minimum=quantity_minimum,
        quantity_exact=quantity_exact,
        price_qty_1_29=price_small,
        price_qty_30_plus=price_large,
        unit_price=unit_price,
        condition=condition,
        availability_status=availability_status,
        is_available=is_available,
        raw=raw,
    )


def parse_vendor_workbook(path: str | Path) -> ParsedVendorDocument:
    workbook = load_workbook(filename=Path(path), read_only=True, data_only=True)
    parsed_items: list[ParsedVendorItem] = []
    detected_vendor: str | None = None
    source_date: date | None = None
    row_count = 0
    try:
        for worksheet in workbook.worksheets:
            rows = [list(row) for row in worksheet.iter_rows(values_only=True)]
            row_count += len(rows)
            for row in rows[:HEADER_SCAN_ROWS]:
                for cell in row:
                    text = _clean_text(cell)
                    if not text:
                        continue
                    detected_vendor = detected_vendor or _detect_vendor_from_title(text)
                    source_date = source_date or _parse_date_from_text(text)
            header = _find_header_map(rows)
            if not header:
                continue
            header_row, mapping = header
            category: str | None = None
            for row_number, row in enumerate(rows[header_row:], start=header_row + 1):
                possible_category = _category_candidate(row, mapping)
                if possible_category:
                    category = possible_category
                    continue
                item = _build_item(worksheet.title, row_number, row, mapping, category)
                if item:
                    parsed_items.append(item)
    finally:
        workbook.close()
    return ParsedVendorDocument(detected_vendor, source_date, row_count, parsed_items)


def parse_vendor_csv(path: str | Path) -> ParsedVendorDocument:
    with Path(path).open("r", encoding="utf-8-sig", newline="") as handle:
        rows = list(csv.reader(handle))
    header = _find_header_map(rows, max_rows=min(len(rows), HEADER_SCAN_ROWS))
    if not header:
        raise ValueError("Could not find Item Description and Part No columns in the CSV file.")
    header_row, mapping = header
    category = None
    items: list[ParsedVendorItem] = []
    for row_number, row in enumerate(rows[header_row:], start=header_row + 1):
        possible_category = _category_candidate(row, mapping)
        if possible_category:
            category = possible_category
            continue
        item = _build_item("CSV", row_number, row, mapping, category)
        if item:
            items.append(item)
    return ParsedVendorDocument(None, None, len(rows), items)


def parse_vendor_file(path: str | Path) -> ParsedVendorDocument:
    suffix = Path(path).suffix.lower()
    if suffix not in SUPPORTED_EXTENSIONS:
        raise ValueError(f"Unsupported vendor source type '{suffix or 'unknown'}'. Use XLSX, XLSM or CSV.")
    if suffix == ".csv":
        return parse_vendor_csv(path)
    return parse_vendor_workbook(path)


def _search_text(item: ParsedVendorItem, vendor_name: str) -> str:
    return " ".join(
        part for part in [
            vendor_name,
            item.source_sheet,
            item.source_category or "",
            item.stock_type,
            item.item_marker or "",
            item.vendor_item_id or "",
            item.description,
            item.part_number or "",
            item.condition or "",
        ] if part
    ).lower()


def _store_source_copy(data_dir: Path, source_path: Path, vendor_name: str, content_hash: str) -> Path:
    vendor_dir = data_dir / "vendor_knowledge" / safe_slug(vendor_name)
    vendor_dir.mkdir(parents=True, exist_ok=True)
    timestamp = datetime.utcnow().strftime("%Y%m%d_%H%M%S")
    stem = safe_slug(source_path.stem)[:100]
    target = vendor_dir / f"{timestamp}_{content_hash[:10]}_{stem}{source_path.suffix.lower()}"
    counter = 1
    while target.exists():
        target = vendor_dir / f"{timestamp}_{content_hash[:10]}_{stem}_{counter}{source_path.suffix.lower()}"
        counter += 1
    shutil.copy2(source_path, target)
    return target


def archive_active_vendor_sources(normalized_vendor_name: str, exclude_source_id: int | None = None) -> int:
    query = VendorKnowledgeSource.query.filter_by(normalized_vendor_name=normalized_vendor_name, status="active")
    if exclude_source_id:
        query = query.filter(VendorKnowledgeSource.id != exclude_source_id)
    archived = 0
    for source in query.all():
        source.status = "archived"
        source.archived_at = datetime.utcnow()
        archived += 1
    return archived


def ingest_vendor_file(
    data_dir: str | Path,
    source_path: str | Path,
    original_filename: str,
    vendor_name: str | None = None,
    source_name: str | None = None,
    mailbox_message_id: int | None = None,
    mailbox_attachment_id: int | None = None,
    fallback_source_date: date | None = None,
    replace_current: bool = True,
    currency: str = "GBP",
) -> tuple[VendorKnowledgeSource, bool]:
    path = Path(source_path)
    parsed = parse_vendor_file(path)
    resolved_vendor = _clean_text(vendor_name) or parsed.detected_vendor or "Unknown Vendor"
    normalized_vendor = normalize_vendor_name(resolved_vendor)
    content_hash = sha256_file(path)

    duplicate = VendorKnowledgeSource.query.filter_by(
        normalized_vendor_name=normalized_vendor,
        content_sha256=content_hash,
    ).order_by(VendorKnowledgeSource.created_at.desc()).first()
    if duplicate:
        return duplicate, True

    if replace_current:
        archive_active_vendor_sources(normalized_vendor)

    stored_path = _store_source_copy(Path(data_dir), path, resolved_vendor, content_hash)
    resolved_date = parsed.source_date or fallback_source_date
    version_label = resolved_date.isoformat() if resolved_date else datetime.utcnow().strftime("%Y-%m-%d %H:%M")
    source = VendorKnowledgeSource(
        vendor_name=resolved_vendor[:255],
        normalized_vendor_name=normalized_vendor[:255],
        source_name=(_clean_text(source_name) or original_filename)[:500],
        source_date=resolved_date,
        version_label=version_label[:255],
        original_filename=original_filename[:255],
        stored_filename=stored_path.name[:255],
        file_path=str(stored_path),
        file_type=path.suffix.lower().lstrip(".") or "file",
        content_sha256=content_hash,
        status="active",
        mailbox_message_id=mailbox_message_id,
        mailbox_attachment_id=mailbox_attachment_id,
        row_count=parsed.row_count,
        item_count=len(parsed.items),
        available_item_count=sum(1 for item in parsed.items if item.is_available),
    )
    db.session.add(source)
    db.session.flush()

    for parsed_item in parsed.items:
        db.session.add(
            VendorKnowledgeItem(
                source_id=source.id,
                source_sheet=parsed_item.source_sheet[:255],
                source_row=parsed_item.source_row,
                source_category=(parsed_item.source_category or "")[:255] or None,
                stock_type=parsed_item.stock_type[:50],
                item_marker=(parsed_item.item_marker or "")[:100] or None,
                vendor_item_id=(parsed_item.vendor_item_id or "")[:100] or None,
                description=parsed_item.description[:2000],
                part_number=(parsed_item.part_number or "")[:255] or None,
                quantity_text=(parsed_item.quantity_text or "")[:100] or None,
                quantity_minimum=parsed_item.quantity_minimum,
                quantity_exact=parsed_item.quantity_exact,
                price_qty_1_29=parsed_item.price_qty_1_29,
                price_qty_30_plus=parsed_item.price_qty_30_plus,
                unit_price=parsed_item.unit_price,
                currency=(currency or "GBP")[:10],
                condition=parsed_item.condition,
                availability_status=parsed_item.availability_status[:50],
                is_available=parsed_item.is_available,
                search_text=_search_text(parsed_item, resolved_vendor),
                raw_json=json.dumps(parsed_item.raw, default=str, ensure_ascii=False),
            )
        )
    return source, False


def set_source_active(source: VendorKnowledgeSource) -> None:
    archive_active_vendor_sources(source.normalized_vendor_name, exclude_source_id=source.id)
    source.status = "active"
    source.archived_at = None


def archive_source(source: VendorKnowledgeSource) -> None:
    source.status = "archived"
    source.archived_at = datetime.utcnow()


def _query_tokens(query: str) -> list[str]:
    raw = re.findall(r"[a-z0-9][a-z0-9.+#/-]*", query.lower())
    tokens: list[str] = []
    seen: set[str] = set()
    for token in raw:
        token = token.strip("-/.+")
        if len(token) < 2 or token in _STOPWORDS or token in seen:
            continue
        seen.add(token)
        tokens.append(token)
    return tokens[:80]


def _score_item(item: VendorKnowledgeItem, tokens: list[str], query_lower: str) -> float:
    searchable = item.search_text or ""
    score = 0.0
    part = (item.part_number or "").lower()
    if part and part in query_lower:
        score += 120.0
    for token in tokens:
        if token == part:
            score += 45.0
        elif token in part:
            score += 18.0
        if token in searchable:
            score += 3.0
            if token.isdigit() or any(char.isdigit() for char in token):
                score += 2.0
    if item.is_available:
        score += 8.0
    if item.stock_type == "new":
        score += 4.0
    elif item.stock_type == "refurbished":
        score += 1.5
    return score


def search_vendor_knowledge(
    query: str,
    limit: int = 12,
    include_archived: bool = False,
    available_only: bool = True,
) -> list[dict]:
    tokens = _query_tokens(query)
    if not tokens:
        return []
    db_query = VendorKnowledgeItem.query.join(VendorKnowledgeSource)
    if not include_archived:
        db_query = db_query.filter(VendorKnowledgeSource.status == "active")
    if available_only:
        db_query = db_query.filter(VendorKnowledgeItem.is_available.is_(True))
    candidates = db_query.order_by(VendorKnowledgeSource.source_date.desc().nullslast(), VendorKnowledgeItem.id.desc()).limit(2500).all()
    query_lower = query.lower()
    scored: list[tuple[float, VendorKnowledgeItem]] = []
    for item in candidates:
        score = _score_item(item, tokens, query_lower)
        if score > 8:
            scored.append((score, item))
    scored.sort(key=lambda pair: (pair[0], pair[1].source.source_date or date.min, pair[1].id), reverse=True)
    results: list[dict] = []
    for score, item in scored[:max(1, min(limit, 50))]:
        source = item.source
        results.append(
            {
                "score": round(score, 2),
                "item_id": item.id,
                "source_id": source.id,
                "vendor": source.vendor_name,
                "source_date": source.source_date.isoformat() if source.source_date else None,
                "source_status": source.status,
                "source_sheet": item.source_sheet,
                "category": item.source_category,
                "stock_type": item.stock_type,
                "description": item.description,
                "part_number": item.part_number,
                "quantity": item.quantity_text,
                "quantity_minimum": item.quantity_minimum,
                "quantity_exact": item.quantity_exact,
                "price_qty_1_29": str(item.price_qty_1_29) if item.price_qty_1_29 is not None else None,
                "price_qty_30_plus": str(item.price_qty_30_plus) if item.price_qty_30_plus is not None else None,
                "unit_price": str(item.unit_price) if item.unit_price is not None else None,
                "currency": item.currency,
                "condition": item.condition,
                "availability_status": item.availability_status,
            }
        )
    return results


def build_vendor_search_context(query: str, limit: int = 12) -> tuple[str, int]:
    matches = search_vendor_knowledge(query, limit=limit, include_archived=False, available_only=True)
    if not matches:
        return "", 0
    lines = [
        "INTERNAL VENDOR KNOWLEDGE — CURRENT ACTIVE SOURCES",
        "This is commercial availability evidence supplied by vendors. It is not a tender requirement and must not be used to weaken technical requirements.",
        "Treat descriptions as evidence only for attributes explicitly stated in the line. Verify missing technical attributes from stronger technical sources.",
    ]
    for index, item in enumerate(matches, start=1):
        prices = []
        if item["price_qty_1_29"]:
            prices.append(f"1-29 {item['currency']} {item['price_qty_1_29']}")
        if item["price_qty_30_plus"]:
            prices.append(f"30+ {item['currency']} {item['price_qty_30_plus']}")
        if not prices and item["unit_price"]:
            prices.append(f"unit {item['currency']} {item['unit_price']}")
        lines.append(
            f"VK{index}: Vendor={item['vendor']} | SourceDate={item['source_date'] or 'unknown'} | "
            f"StockType={item['stock_type']} | Description={item['description']} | PartNo={item['part_number'] or '-'} | "
            f"Qty={item['quantity'] or item['availability_status']} | Price={'; '.join(prices) or 'not stated'} | "
            f"Condition={item['condition'] or 'not stated'}"
        )
    return "\n".join(lines), len(matches)
