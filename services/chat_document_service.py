from __future__ import annotations

import re
from io import BytesIO

from docx import Document


def build_chat_document(title: str, markdown_text: str) -> bytes:
    """Create a straightforward editable DOCX from an AI-drafted Markdown document."""
    document = Document()
    document.core_properties.title = title
    lines = (markdown_text or "").replace("\r\n", "\n").split("\n")
    for raw in lines:
        line = raw.rstrip()
        if not line:
            document.add_paragraph()
            continue
        heading = re.match(r"^(#{1,4})\s+(.+)$", line)
        if heading:
            document.add_heading(heading.group(2).strip(), level=min(len(heading.group(1)), 4))
            continue
        bullet = re.match(r"^[-*+]\s+(.+)$", line)
        if bullet:
            document.add_paragraph(bullet.group(1).strip(), style="List Bullet")
            continue
        numbered = re.match(r"^\d+[.)]\s+(.+)$", line)
        if numbered:
            document.add_paragraph(numbered.group(1).strip(), style="List Number")
            continue
        document.add_paragraph(line)
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def safe_document_filename(title: str, extension: str = ".docx") -> str:
    stem = re.sub(r"[^A-Za-z0-9._-]+", "_", title or "Tender_Document").strip("._")[:100]
    return (stem or "Tender_Document") + extension
